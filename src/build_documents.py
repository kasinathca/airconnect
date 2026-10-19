"""Generate the editable academic Word report from the verified Markdown master."""
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "PROJECT_REPORT.md"
OUTPUT = ROOT / "docs" / "AirConnect_India_Project_Report.docx"
CHARTS = ROOT / "docs" / "charts"
NAVY, LIGHT, BORDER = "173B58", "EAF1F6", "D5DDE3"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc; tc_pr = tc.get_or_add_tcPr(); margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None: margins = OxmlElement("w:tcMar"); tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None: node = OxmlElement(f"w:{name}"); margins.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")


def rich(paragraph, text):
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text)
    for part in parts:
        if not part: continue
        if part.startswith("**") and part.endswith("**"): run = paragraph.add_run(part[2:-2]); run.bold = True
        elif part.startswith("*") and part.endswith("*"): run = paragraph.add_run(part[1:-1]); run.italic = True
        elif part.startswith("`") and part.endswith("`"): run = paragraph.add_run(part[1:-1]); run.font.name = "Consolas"; run.font.size = Pt(9)
        else: paragraph.add_run(part)


def add_figure(doc, filename, caption, width=6.35):
    path = CHARTS / filename
    if not path.exists(): return
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption); run.italic = True; run.font.size = Pt(9); run.font.color.rgb = RGBColor(74, 88, 99)


def add_workflow(doc):
    labels = ["Pinned source", "Validation", "Exact lags", "Chronological models", "Canonical JSON", "Website and documents"]
    table = doc.add_table(rows=1, cols=len(labels)); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
    for index, (cell, label) in enumerate(zip(table.rows[0].cells, labels)):
        cell.width = Inches(1.05); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cell, 100, 70, 100, 70)
        set_cell_shading(cell, LIGHT if index % 2 == 0 else "F6F8FA")
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run(label); r.bold = True; r.font.size = Pt(8); r.font.color.rgb = RGBColor(23, 59, 88)
    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER; run = cap.add_run("Figure 1. Reproducible AirConnect India v2.0 system workflow."); run.italic = True; run.font.size = Pt(9)


doc = Document()
section = doc.sections[0]; section.top_margin = Inches(.72); section.bottom_margin = Inches(.7); section.left_margin = Inches(.86); section.right_margin = Inches(.78)
styles = doc.styles
normal = styles["Normal"]; normal.font.name = "Times New Roman"; normal.font.size = Pt(10.5); normal.font.color.rgb = RGBColor(25, 25, 25); normal.paragraph_format.line_spacing = 1.14; normal.paragraph_format.space_after = Pt(5)
for name, size in [("Title", 20), ("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)]:
    style = styles[name]; style.font.name = "Times New Roman"; style.font.size = Pt(size); style.font.bold = True; style.font.color.rgb = RGBColor(0, 0, 0); style.paragraph_format.space_before = Pt(10); style.paragraph_format.space_after = Pt(5); style.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT; hr = header.add_run("AirConnect India v2.0  |  MDI4001"); hr.font.name = "Times New Roman"; hr.font.size = Pt(8); hr.font.color.rgb = RGBColor(90, 100, 108)
footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER; fr = footer.add_run("Vellore Institute of Technology  |  Machine Learning for Data Science"); fr.font.name = "Times New Roman"; fr.font.size = Pt(8); fr.font.color.rgb = RGBColor(90, 100, 108)

figures = {
    "3. Related work and data provenance": ("annual_traffic.png", "Figure 2. National domestic and international airport passenger movements, 2016–2025."),
    "10. Regression results": ("regression_model_comparison.png", "Figure 3. Held-out 2025 regression MAE for trained models and baselines."),
    "11. Regression error analysis": ("regression_actual_predicted.png", "Figure 4. Actual and predicted 2025 airport-month movements for the validation-selected Random Forest."),
    "13. Classification results": ("classification_confusion_matrix.png", "Figure 5. Logistic Regression confusion matrix on 1,357 untouched 2025 observations."),
    "14. K-Means clustering": ("clustering_selection.png", "Figure 6. K-Means inertia and Silhouette Score evidence for k = 2–6."),
    "15. Principal component analysis": ("pca_clusters.png", "Figure 7. PCA projection of the six-feature airport matrix, coloured by K-Means group."),
}

lines = SOURCE.read_text(encoding="utf-8").splitlines(); i = 0; in_code = False; code_lines = []; current_title = ""
while i < len(lines):
    raw = lines[i]; line = raw.strip()
    if line.startswith("```"):
        if in_code:
            if current_title != "5. System architecture":
                p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(.3); p.paragraph_format.right_indent = Inches(.3)
                run = p.add_run("\n".join(code_lines)); run.font.name = "Consolas"; run.font.size = Pt(8)
            code_lines = []; in_code = False
        else: in_code = True
        i += 1; continue
    if in_code: code_lines.append(raw); i += 1; continue
    if not line or line == "---": i += 1; continue
    if line.startswith("#"):
        level = len(line) - len(line.lstrip("#")); title = line[level:].strip()
        current_title = title
        style = "Title" if level == 1 else "Heading 1" if level == 2 else "Heading 2"
        p = doc.add_paragraph(style=style); rich(p, title)
        if title == "5. System architecture": add_workflow(doc)
        if title in figures: add_figure(doc, *figures[title])
        i += 1; continue
    if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")]); i += 1
        rows = [row for row in rows if not all(re.fullmatch(r":?-{2,}:?", cell) for cell in row)]
        cols = max(map(len, rows)); table = doc.add_table(rows=0, cols=cols); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.style = "Table Grid"
        for r_index, values in enumerate(rows):
            cells = table.add_row().cells
            for c_index in range(cols):
                value = values[c_index] if c_index < len(values) else ""; cell = cells[c_index]; cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cell)
                if r_index == 0: set_cell_shading(cell, NAVY)
                elif r_index % 2 == 0: set_cell_shading(cell, "F4F7F9")
                p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_index == 0 else WD_ALIGN_PARAGRAPH.CENTER; rich(p, value)
                for run in p.runs: run.font.name = "Times New Roman"; run.font.size = Pt(8); run.font.color.rgb = RGBColor(255,255,255) if r_index == 0 else RGBColor(25,25,25); run.bold = r_index == 0 or run.bold
        doc.add_paragraph().paragraph_format.space_after = Pt(0); continue
    if line.startswith("- "):
        p = doc.add_paragraph(style="List Bullet"); rich(p, line[2:]); i += 1; continue
    p = doc.add_paragraph(); rich(p, line); i += 1

doc.core_properties.title = "AirConnect India Machine Learning Project Report"
doc.core_properties.subject = "MDI4001 Machine Learning for Data Science"
doc.core_properties.author = "AirConnect India project team"
doc.save(OUTPUT)
print(f"DOCX: {OUTPUT} ({OUTPUT.stat().st_size:,} bytes)")
